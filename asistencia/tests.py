import os
import shutil
import sys
import tempfile
import types
import zipfile
from datetime import date, datetime, time, timedelta
from io import BytesIO
from unittest import mock

from django.conf import settings
from django.contrib.auth.models import Permission, User
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, TransactionTestCase, Client, override_settings
from django.urls import reverse
from django.utils import timezone
from PIL import Image

from .models import Asistencia, ConfiguracionSistema, Evento, TipoEvento, HistorialCarnet, Justificacion, LogAccion, SolicitudDescargaCarnets, Usuario
from .services import AsistenciaService
from .forms import CargaMasivaFotosForm, FiltroAsistenciaForm
from .utils.reports import generate_attendance_csv, get_filtered_attendance_data


class AttendanceEnhancementsTests(TestCase):
    def setUp(self):
        self.temp_media_root = tempfile.mkdtemp(prefix='test-media-')
        self._media_override = override_settings(MEDIA_ROOT=self.temp_media_root)
        self._media_override.enable()

        self.user = User.objects.create_user(username='manager', password='password')
        self.user.user_permissions.add(
            Permission.objects.get(codename='can_manage_users'),
            Permission.objects.get(codename='can_scan_qr'),
        )
        self.client = Client()
        self.client.force_login(self.user)

        ConfiguracionSistema.objects.create(nombre_institucion='QUIULACOCHA', tolerancia_minutos=15)
        self.evento = Evento.objects.create(
            nombre='Asamblea General',
            fecha=date.today(),
            hora_ingreso=time(8, 0),
            activo=True,
        )

    def tearDown(self):
        self._media_override.disable()
        shutil.rmtree(self.temp_media_root, ignore_errors=True)

    def test_registrar_ingreso_marca_puntualidad(self):
        socio_puntual = Usuario.objects.create(nombre='Ana', apellido='Perez', dni='12345670', estado=Usuario.ESTADO_ACTIVO)
        socio_tarde = Usuario.objects.create(nombre='Luis', apellido='Rojas', dni='12345671', estado=Usuario.ESTADO_ACTIVO)

        ts_puntual = timezone.make_aware(datetime.combine(date.today(), time(8, 10)))
        ts_tarde = timezone.make_aware(datetime.combine(date.today(), time(8, 30)))

        a1, _, _ = AsistenciaService.registrar_ingreso(
            usuario=socio_puntual,
            evento=self.evento,
            ubicacion=None,
            request_user=self.user,
            fecha_registro=ts_puntual,
        )
        a2, _, _ = AsistenciaService.registrar_ingreso(
            usuario=socio_tarde,
            evento=self.evento,
            ubicacion=None,
            request_user=self.user,
            fecha_registro=ts_tarde,
        )

        self.assertEqual(a1.puntualidad, Asistencia.PUNTUALIDAD_PUNTUAL)
        self.assertEqual(a2.puntualidad, Asistencia.PUNTUALIDAD_TARDE)

    def test_tardanza_activa_convierte_tarde_en_falta_en_reportes(self):
        socio_tarde = Usuario.objects.create(nombre='Luis', apellido='Rojas', dni='12345671', estado=Usuario.ESTADO_ACTIVO)
        ts_tarde = timezone.make_aware(datetime.combine(date.today(), time(8, 30)))

        AsistenciaService.registrar_ingreso(
            usuario=socio_tarde,
            evento=self.evento,
            ubicacion=None,
            request_user=self.user,
            fecha_registro=ts_tarde,
        )

        data = get_filtered_attendance_data({'evento': self.evento})
        unified_list = data['unified_report']

        self.assertEqual(len(unified_list), 1)
        self.assertTrue(unified_list[0].is_absent)
        self.assertTrue(unified_list[0].is_late_absent)

        csv_response = generate_attendance_csv(unified_list)
        csv_content = csv_response.content.decode('utf-8')

        self.assertIn('FALTA', csv_content)
        self.assertIn('Tardanza', csv_content)

    def test_tardanza_desactivada_mantiene_registro_como_presente(self):
        config = ConfiguracionSistema.objects.get()
        config.tardanza_activa = False
        config.save(update_fields=['tardanza_activa'])

        socio_tarde = Usuario.objects.create(nombre='Luis', apellido='Rojas', dni='12345671', estado=Usuario.ESTADO_ACTIVO)
        ts_tarde = timezone.make_aware(datetime.combine(date.today(), time(8, 30)))

        AsistenciaService.registrar_ingreso(
            usuario=socio_tarde,
            evento=self.evento,
            ubicacion=None,
            request_user=self.user,
            fecha_registro=ts_tarde,
        )

        data = get_filtered_attendance_data({'evento': self.evento})
        unified_list = data['unified_report']

        self.assertEqual(len(unified_list), 1)
        self.assertFalse(unified_list[0].is_absent)
        self.assertFalse(getattr(unified_list[0], 'is_late_absent', False))
        self.assertEqual(unified_list[0].puntualidad, Asistencia.PUNTUALIDAD_TARDE)

    def test_exonerado_tarde_no_se_convierte_en_falta(self):
        socio_exonerado = Usuario.objects.create(
            nombre='Julia',
            apellido='Diaz',
            dni='12345699',
            estado=Usuario.ESTADO_EXONERADO,
        )
        Asistencia.objects.create(
            usuario=socio_exonerado,
            evento=self.evento,
            fecha=self.evento.fecha,
            hora_ingreso=time(8, 30),
            confirmada=True,
            puntualidad=Asistencia.PUNTUALIDAD_TARDE,
        )

        data = get_filtered_attendance_data({'evento': self.evento})
        unified_list = data['unified_report']

        self.assertEqual(len(unified_list), 1)
        self.assertFalse(unified_list[0].is_absent)
        self.assertFalse(getattr(unified_list[0], 'is_late_absent', False))

    def test_filtro_faltaron_ignora_confirmada_para_ausencias(self):
        socio_tarde = Usuario.objects.create(nombre='Luis', apellido='Rojas', dni='12345671', estado=Usuario.ESTADO_ACTIVO)
        socio_justificado = Usuario.objects.create(nombre='Marta', apellido='Lopez', dni='12345672', estado=Usuario.ESTADO_ACTIVO)
        socio_sin_registro = Usuario.objects.create(nombre='Jose', apellido='Diaz', dni='12345673', estado=Usuario.ESTADO_ACTIVO)

        ts_tarde = timezone.make_aware(datetime.combine(date.today(), time(8, 30)))
        AsistenciaService.registrar_ingreso(
            usuario=socio_tarde,
            evento=self.evento,
            ubicacion=None,
            request_user=self.user,
            fecha_registro=ts_tarde,
        )

        Asistencia.objects.create(
            usuario=socio_justificado,
            evento=self.evento,
            fecha=self.evento.fecha,
            confirmada=True,
            es_justificada=True,
        )

        data_sin_confirmada = get_filtered_attendance_data({'evento': self.evento, 'estado': 'faltaron'})
        data_confirmada_true = get_filtered_attendance_data({'evento': self.evento, 'estado': 'faltaron', 'confirmada': 'true'})
        data_confirmada_false = get_filtered_attendance_data({'evento': self.evento, 'estado': 'faltaron', 'confirmada': 'false'})

        def resumen(data):
            return sorted(
                (
                    item.usuario.dni if hasattr(item, 'usuario') else item.dni,
                    bool(getattr(item, 'is_late_absent', False)),
                    bool(getattr(item, 'es_justificada', False)),
                )
                for item in data['unified_report']
            )

        esperado = resumen(data_sin_confirmada)
        self.assertIn((socio_tarde.dni, True, False), esperado)
        self.assertIn((socio_justificado.dni, False, True), esperado)
        self.assertIn((socio_sin_registro.dni, False, False), esperado)
        self.assertEqual(esperado, resumen(data_confirmada_true))
        self.assertEqual(esperado, resumen(data_confirmada_false))

    def test_cerrar_evento_lo_inactiva_y_genera_resumen(self):
        response = self.client.post(reverse('cerrar_evento', args=[self.evento.id]))
        self.assertEqual(response.status_code, 200)

        self.evento.refresh_from_db()
        self.assertFalse(self.evento.activo)

        payload = response.json()
        self.assertIn('pdf_url', payload)
        self.assertIn('excel_url', payload)

        pdf_name = os.path.basename(payload['pdf_url'])
        excel_name = os.path.basename(payload['excel_url'])
        self.assertTrue(os.path.exists(os.path.join(settings.MEDIA_ROOT, 'cierres_evento', pdf_name)))
        self.assertTrue(os.path.exists(os.path.join(settings.MEDIA_ROOT, 'cierres_evento', excel_name)))

        audit_log = LogAccion.objects.filter(accion__icontains='Auditoria').order_by('-fecha').first()
        self.assertIsNotNone(audit_log)

    def test_cerrar_evento_bloquea_si_hay_inconsistencia_exonerados(self):
        Usuario.objects.create(
            nombre='Exo',
            apellido='SinRegistro',
            dni='12345999',
            estado=Usuario.ESTADO_EXONERADO,
        )

        response = self.client.post(reverse('cerrar_evento', args=[self.evento.id]))
        self.assertEqual(response.status_code, 409)
        payload = response.json()
        self.assertEqual(payload.get('code'), 'EXONERADO_MISMATCH')

        self.evento.refresh_from_db()
        self.assertTrue(self.evento.activo)

    def test_actualizar_foto_rapida_sube_imagen_y_devuelve_fragmento_htmx(self):
        usuario = Usuario.objects.create(nombre='Elena', apellido='Quispe', dni='12345674', estado=Usuario.ESTADO_ACTIVO)
        source = BytesIO()
        Image.new('RGB', (1200, 900), color=(24, 120, 196)).save(source, format='PNG')
        imagen = SimpleUploadedFile('avatar.png', source.getvalue(), content_type='image/png')

        response = self.client.post(
            reverse('actualizar_foto_rapida', args=[usuario.dni]),
            {'foto': imagen},
            HTTP_HX_REQUEST='true',
        )

        self.assertEqual(response.status_code, 200)
        usuario.refresh_from_db()
        self.assertTrue(bool(usuario.foto_perfil))
        self.assertTrue(usuario.foto_perfil.name.endswith('.jpg'))

        with Image.open(usuario.foto_perfil.path) as stored_image:
            self.assertEqual(stored_image.size, Usuario.FOTO_PERFIL_SIZE)
            self.assertEqual(stored_image.format, 'JPEG')

        content = response.content.decode()
        self.assertIn(f'id="avatar-container-{usuario.dni}"', content)
        self.assertIn('hx-trigger="change"', content)
        self.assertIn('bi-camera-fill', content)
        self.assertTrue(
            LogAccion.objects.filter(
                usuario=self.user,
                accion='Actualización de foto rápida',
            ).exists()
        )

    def test_carga_masiva_form_acepta_varias_imagenes(self):
        first_buffer = BytesIO()
        second_buffer = BytesIO()
        Image.new('RGB', (200, 200), color=(20, 120, 180)).save(first_buffer, format='JPEG')
        Image.new('RGB', (200, 200), color=(180, 80, 40)).save(second_buffer, format='PNG')

        foto_1 = SimpleUploadedFile('12345678.jpg', first_buffer.getvalue(), content_type='image/jpeg')
        foto_2 = SimpleUploadedFile('87654321.png', second_buffer.getvalue(), content_type='image/png')

        form = CargaMasivaFotosForm(files={'fotos': [foto_1, foto_2]})

        self.assertTrue(form.is_valid(), form.errors.as_json())
        self.assertEqual(len(form.cleaned_data['fotos']), 2)

    def test_actualizar_foto_rapida_requiere_permiso(self):
        usuario = Usuario.objects.create(nombre='Mario', apellido='Lopez', dni='12345675', estado=Usuario.ESTADO_ACTIVO)
        no_manager = User.objects.create_user(username='viewer', password='password')
        self.client.force_login(no_manager)

        response = self.client.post(reverse('actualizar_foto_rapida', args=[usuario.dni]))

        self.assertEqual(response.status_code, 403)

    def test_filtro_asistencia_rechaza_rango_invalido(self):
        form = FiltroAsistenciaForm(data={
            'fecha_inicio': '2026-04-10',
            'fecha_fin': '2026-04-01',
        })

        self.assertFalse(form.is_valid())
        self.assertIn('La fecha inicio no puede ser mayor que la fecha fin.', form.errors['__all__'])

    def test_exportar_excel_respeta_filtro_confirmada(self):
        usuario_confirmado = Usuario.objects.create(nombre='Rosa', apellido='Diaz', dni='12345676', estado=Usuario.ESTADO_ACTIVO)
        usuario_pendiente = Usuario.objects.create(nombre='Saul', apellido='Perez', dni='12345677', estado=Usuario.ESTADO_ACTIVO)

        Asistencia.objects.create(
            usuario=usuario_confirmado,
            fecha=self.evento.fecha,
            hora_ingreso=time(8, 5),
            hora_salida=time(10, 0),
            evento=self.evento,
            confirmada=True,
        )
        Asistencia.objects.create(
            usuario=usuario_pendiente,
            fecha=self.evento.fecha,
            hora_ingreso=time(8, 10),
            evento=self.evento,
            confirmada=False,
        )

        response = self.client.get(reverse('exportar_asistencias_excel'), {
            'evento': self.evento.id,
            'confirmada': 'true',
        })

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response['Content-Type']
        )

    def test_exportar_excel_diversificado_con_pestanas_y_faltantes(self):
        import openpyxl
        socio_asistio = Usuario.objects.create(nombre='Ana', apellido='Gomez', dni='88990011', estado=Usuario.ESTADO_ACTIVO)
        socio_falto = Usuario.objects.create(nombre='Pedro', apellido='Rios', dni='88990022', estado=Usuario.ESTADO_ACTIVO)
        socio_justificado = Usuario.objects.create(nombre='Maria', apellido='Salas', dni='88990033', estado=Usuario.ESTADO_ACTIVO)

        Asistencia.objects.create(
            usuario=socio_asistio,
            fecha=self.evento.fecha,
            hora_ingreso=time(8, 0),
            hora_salida=time(10, 0),
            evento=self.evento,
            confirmada=True,
        )
        Justificacion.objects.create(
            usuario=socio_justificado,
            evento=self.evento,
            motivo='Cita medica',
            estado=Justificacion.ESTADO_APROBADO,
        )

        response = self.client.get(reverse('exportar_asistencias_excel'), {
            'evento': self.evento.id,
        })
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertIn('Resumen Ejecutivo', wb.sheetnames)
        self.assertIn('Faltas Reales', wb.sheetnames)
        self.assertIn('Asistieron', wb.sheetnames)
        self.assertIn('Justificaciones', wb.sheetnames)
        self.assertIn('Pendientes', wb.sheetnames)
        self.assertIn('Exonerados', wb.sheetnames)
        self.assertIn('Padrón General', wb.sheetnames)

        ws_faltas = wb['Faltas Reales']
        self.assertIsNotNone(ws_faltas.auto_filter.ref)

        # Probar también exportación CSV con inasistentes
        response_csv = self.client.get(reverse('exportar_asistencias_csv'), {
            'evento': self.evento.id,
        })
        self.assertEqual(response_csv.status_code, 200)
        self.assertIn('text/csv', response_csv['Content-Type'])

    def test_exportar_excel_filtrando_solo_por_fecha_inicio_incluye_padron_completo(self):
        import openpyxl
        socio_asistio = Usuario.objects.create(nombre='Luis', apellido='Perez', dni='77889901', estado=Usuario.ESTADO_ACTIVO)
        socio_falto = Usuario.objects.create(nombre='Rosa', apellido='Castro', dni='77889902', estado=Usuario.ESTADO_ACTIVO)

        Asistencia.objects.create(
            usuario=socio_asistio,
            fecha=self.evento.fecha,
            hora_ingreso=time(8, 0),
            hora_salida=time(10, 0),
            evento=self.evento,
            confirmada=True,
        )

        # Filtrar solo por fecha_inicio (sin fecha_fin y sin pasar evento)
        response = self.client.get(reverse('exportar_asistencias_excel'), {
            'fecha_inicio': self.evento.fecha.strftime('%Y-%m-%d'),
        })
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertIn('Padrón General', wb.sheetnames)
        self.assertIn('Faltas Reales', wb.sheetnames)

        # Debe incluir a los inasistentes (socio_falto) en Faltas Reales y Padrón General
        ws_general = wb['Padrón General']
        self.assertGreater(ws_general.max_row, 5) # Encabezados + datos
        ws_faltas = wb['Faltas Reales']
        self.assertGreater(ws_faltas.max_row, 5) # Debe contener faltantes

    def test_exportar_reporte_global_excel_incluye_matriz_y_consolidado(self):
        import openpyxl
        socio_1 = Usuario.objects.create(nombre='Carlos', apellido='Bendezu', dni='11223344', estado=Usuario.ESTADO_ACTIVO)
        socio_2 = Usuario.objects.create(nombre='Diana', apellido='Zevallos', dni='55667788', estado=Usuario.ESTADO_ACTIVO)

        Asistencia.objects.create(
            usuario=socio_1,
            fecha=self.evento.fecha,
            hora_ingreso=time(8, 0),
            hora_salida=time(10, 0),
            evento=self.evento,
            confirmada=True,
        )

        response = self.client.get(reverse('exportar_reporte_global_excel'))
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertIn('Resumen Ejecutivo', wb.sheetnames)
        self.assertIn('Matriz de Asambleas', wb.sheetnames)
        self.assertIn('Consolidado por Comunero', wb.sheetnames)
        self.assertIn('Inasistencias Acumuladas', wb.sheetnames)
        self.assertIn('Justificaciones Acumuladas', wb.sheetnames)
        self.assertIn('Detalle General', wb.sheetnames)

        ws_matrix = wb['Matriz de Asambleas']
        self.assertEqual(ws_matrix.freeze_panes, "E6")
        self.assertIsNotNone(ws_matrix.auto_filter.ref)

        ws_consolidado = wb['Consolidado por Comunero']
        self.assertIsNotNone(ws_consolidado.auto_filter.ref)

    def test_descargar_reporte_usuario_excel(self):
        import openpyxl
        socio = Usuario.objects.create(nombre='Manuel', apellido='Quispe', dni='99887766', estado=Usuario.ESTADO_ACTIVO)
        Asistencia.objects.create(
            usuario=socio,
            fecha=self.evento.fecha,
            hora_ingreso=time(8, 0),
            hora_salida=time(10, 0),
            evento=self.evento,
            confirmada=True,
        )

        response = self.client.get(reverse('descargar_reporte_usuario_excel', kwargs={'dni': socio.dni}))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertIn('Historial Individual', wb.sheetnames)
        ws = wb['Historial Individual']
        self.assertEqual(ws.freeze_panes, "A11")
        self.assertIsNotNone(ws.auto_filter.ref)

    def test_reporte_usuario_seguridad_sin_fuga_de_datos(self):
        """Verifica que un comunero regular no pueda ver reportes de otros pero sí el suyo."""
        from django.contrib.auth.models import User
        user_a = User.objects.create_user(username='comunero_a', password='password123')
        socio_a = Usuario.objects.create(user=user_a, nombre='Comunero', apellido='A', dni='11112222', estado=Usuario.ESTADO_ACTIVO)

        socio_b = Usuario.objects.create(nombre='Comunero', apellido='B', dni='33334444', estado=Usuario.ESTADO_ACTIVO)

        self.client.force_login(user_a)

        # Intento de acceder a reporte de socio_b -> 403 Forbidden (sin fuga de datos)
        resp_pdf_b = self.client.get(reverse('descargar_reporte_usuario_pdf', kwargs={'dni': socio_b.dni}))
        self.assertEqual(resp_pdf_b.status_code, 403)

        resp_excel_b = self.client.get(reverse('descargar_reporte_usuario_excel', kwargs={'dni': socio_b.dni}))
        self.assertEqual(resp_excel_b.status_code, 403)

        # Acceso a su propio reporte -> 200 OK
        resp_excel_a = self.client.get(reverse('descargar_reporte_usuario_excel', kwargs={'dni': socio_a.dni}))
        self.assertEqual(resp_excel_a.status_code, 200)

    def test_tipo_evento_clasificacion_asamblea_y_faena(self):
        tipo_asamblea = TipoEvento.objects.create(
            nombre='Asamblea Ordinaria Anual',
            categoria=TipoEvento.CATEGORIA_ASAMBLEA,
            color='#1e40af'
        )
        tipo_faena = TipoEvento.objects.create(
            nombre='Faena de Limpieza Comunal',
            categoria=TipoEvento.CATEGORIA_FAENA,
            color='#047857'
        )

        ev_asamblea = Evento.objects.create(
            nombre='Asamblea de Presentación',
            tipo_evento=tipo_asamblea,
            fecha=date(2026, 5, 10),
            hora_ingreso=time(8, 0),
        )
        ev_faena = Evento.objects.create(
            nombre='Faena en Carretera',
            tipo_evento=tipo_faena,
            fecha=date(2026, 6, 12),
            hora_ingreso=time(7, 0),
        )
        ev_historico_faena = Evento.objects.create(
            nombre='FAENA POR FIESTAS PATRIAS',
            fecha=date(2026, 7, 26),
            hora_ingreso=time(13, 0),
        )

        self.assertTrue(ev_asamblea.es_asamblea)
        self.assertFalse(ev_asamblea.es_faena)
        self.assertEqual(ev_asamblea.tipo_nombre, 'Asamblea Ordinaria Anual')

        self.assertTrue(ev_faena.es_faena)
        self.assertFalse(ev_faena.es_asamblea)
        self.assertEqual(ev_faena.tipo_nombre, 'Faena de Limpieza Comunal')

        # Auto-detección histórica por nombre
        self.assertTrue(ev_historico_faena.es_faena)
        self.assertFalse(ev_historico_faena.es_asamblea)
        self.assertEqual(ev_historico_faena.tipo_nombre, 'Faena Comunal')

    def test_filtrar_asistencias_por_faena_y_asamblea(self):
        tipo_faena, _ = TipoEvento.objects.get_or_create(
            nombre='Faena General',
            defaults={'categoria': TipoEvento.CATEGORIA_FAENA, 'color': '#047857'}
        )
        tipo_asamblea, _ = TipoEvento.objects.get_or_create(
            nombre='Asamblea Extraordinaria Comunal',
            defaults={'categoria': TipoEvento.CATEGORIA_ASAMBLEA, 'color': '#7c3aed'}
        )

        ev_faena = Evento.objects.create(
            nombre='Faena Mantenimiento',
            tipo_evento=tipo_faena,
            fecha=date(2026, 8, 10),
            hora_ingreso=time(8, 0),
        )
        ev_asamblea = Evento.objects.create(
            nombre='Asamblea de Presupuesto',
            tipo_evento=tipo_asamblea,
            fecha=date(2026, 8, 15),
            hora_ingreso=time(9, 0),
        )

        socio_1 = Usuario.objects.create(nombre='Javier', apellido='Mayhua', dni='55443322', estado=Usuario.ESTADO_ACTIVO)
        Asistencia.objects.create(
            usuario=socio_1,
            fecha=ev_faena.fecha,
            hora_ingreso=time(8, 5),
            hora_salida=time(12, 0),
            evento=ev_faena,
            confirmada=True,
        )

        # Filtro por FAENA
        data_faena = get_filtered_attendance_data({'categoria_evento': 'FAENA'})
        eventos_faena = [it.evento.id for it in data_faena['unified_report'] if getattr(it, 'evento', None)]
        self.assertTrue(all(ev_id == ev_faena.id for ev_id in eventos_faena if ev_id))

        # Filtro por ASAMBLEA
        data_asamblea = get_filtered_attendance_data({'categoria_evento': 'ASAMBLEA'})
        eventos_asamblea = [it.evento.id for it in data_asamblea['unified_report'] if getattr(it, 'evento', None)]
        self.assertTrue(all(ev_id != ev_faena.id for ev_id in eventos_asamblea if ev_id))

    def test_excel_diferencia_faltas_faenas_vs_asambleas(self):
        import openpyxl
        # Limpiar eventos previos para aislar exactamente los dos eventos de la prueba
        Evento.objects.all().delete()

        tipo_faena, _ = TipoEvento.objects.get_or_create(
            nombre='Faena de Obras',
            defaults={'categoria': TipoEvento.CATEGORIA_FAENA, 'color': '#047857'}
        )
        tipo_asamblea, _ = TipoEvento.objects.get_or_create(
            nombre='Asamblea del Pueblo',
            defaults={'categoria': TipoEvento.CATEGORIA_ASAMBLEA, 'color': '#1e40af'}
        )

        ev_faena = Evento.objects.create(
            nombre='Faena Limpieza Reservorio',
            tipo_evento=tipo_faena,
            fecha=date(2026, 9, 5),
            hora_ingreso=time(7, 30),
        )
        ev_asamblea = Evento.objects.create(
            nombre='Asamblea Elección Directiva',
            tipo_evento=tipo_asamblea,
            fecha=date(2026, 9, 20),
            hora_ingreso=time(8, 30),
        )

        socio_asiste_todo = Usuario.objects.create(nombre='Héctor', apellido='Alanya', dni='10203040', estado=Usuario.ESTADO_ACTIVO)
        socio_falta_faena = Usuario.objects.create(nombre='Marcos', apellido='Casas', dni='20304050', estado=Usuario.ESTADO_ACTIVO)
        socio_falta_todo = Usuario.objects.create(nombre='Nelly', apellido='Meza', dni='30405060', estado=Usuario.ESTADO_ACTIVO)

        # Héctor asiste a ambos
        Asistencia.objects.create(usuario=socio_asiste_todo, evento=ev_faena, fecha=ev_faena.fecha, hora_ingreso=time(7, 30), hora_salida=time(12, 0), confirmada=True)
        Asistencia.objects.create(usuario=socio_asiste_todo, evento=ev_asamblea, fecha=ev_asamblea.fecha, hora_ingreso=time(8, 30), hora_salida=time(13, 0), confirmada=True)

        # Marcos asiste solo a asamblea (falta a faena)
        Asistencia.objects.create(usuario=socio_falta_faena, evento=ev_asamblea, fecha=ev_asamblea.fecha, hora_ingreso=time(8, 30), hora_salida=time(13, 0), confirmada=True)

        # Nelly no asiste a ninguno (falta a faena y asamblea)

        # 1. Probar exportar_asistencias_excel
        response_excel = self.client.get(reverse('exportar_asistencias_excel'))
        self.assertEqual(response_excel.status_code, 200)
        wb = openpyxl.load_workbook(BytesIO(response_excel.content))

        # Verificar que se generaron las pestañas especializadas de faenas y asambleas
        self.assertIn('Faltas a Faenas', wb.sheetnames)
        self.assertIn('Faltas a Asambleas', wb.sheetnames)
        self.assertIn('Consolidado por Comunero', wb.sheetnames)

        ws_faenas = wb['Faltas a Faenas']
        ws_asambleas = wb['Faltas a Asambleas']
        ws_consolidado = wb['Consolidado por Comunero']

        # En Faltas a Faenas deben estar Marcos y Nelly (no Héctor)
        dnis_faltas_faenas = [str(ws_faenas.cell(row=r, column=3).value) for r in range(6, ws_faenas.max_row + 1)]
        self.assertIn('20304050', dnis_faltas_faenas)
        self.assertIn('30405060', dnis_faltas_faenas)
        self.assertNotIn('10203040', dnis_faltas_faenas)

        # En Faltas a Asambleas debe estar Nelly (no Marcos ni Héctor)
        dnis_faltas_asambleas = [str(ws_asambleas.cell(row=r, column=3).value) for r in range(6, ws_asambleas.max_row + 1)]
        self.assertIn('30405060', dnis_faltas_asambleas)
        self.assertNotIn('20304050', dnis_faltas_asambleas)
        self.assertNotIn('10203040', dnis_faltas_asambleas)

        # 2. Probar exportar_reporte_global_excel
        response_global = self.client.get(reverse('exportar_reporte_global_excel'))
        self.assertEqual(response_global.status_code, 200)
        wb_global = openpyxl.load_workbook(BytesIO(response_global.content))

        self.assertIn('Matriz de Asambleas', wb_global.sheetnames)
        self.assertIn('Matriz de Faenas', wb_global.sheetnames)
        self.assertIn('Sábana General', wb_global.sheetnames)
        self.assertIn('Consolidado por Comunero', wb_global.sheetnames)
        self.assertIn('Faltas a Faenas', wb_global.sheetnames)
        self.assertIn('Faltas a Asambleas', wb_global.sheetnames)

        # Probar filtro categoria=FAENA en exportar_reporte_global_excel
        response_solo_faenas = self.client.get(reverse('exportar_reporte_global_excel'), {'categoria': 'FAENA'})
        self.assertEqual(response_solo_faenas.status_code, 200)
        wb_faenas_global = openpyxl.load_workbook(BytesIO(response_solo_faenas.content))
        self.assertIn('Consolidado por Comunero', wb_faenas_global.sheetnames)

    def test_historial_con_filtro_avanzado_por_fecha_muestra_resultados(self):
        evento_marzo = Evento.objects.create(
            nombre='Reunion Marzo',
            fecha=date(2026, 3, 15),
            hora_ingreso=time(8, 0),
            activo=False,
        )
        evento_abril = Evento.objects.create(
            nombre='Reunion Abril',
            fecha=date(2026, 4, 2),
            hora_ingreso=time(8, 0),
            activo=False,
        )
        usuario = Usuario.objects.create(nombre='Julia', apellido='Soto', dni='12345679', estado=Usuario.ESTADO_ACTIVO)

        Asistencia.objects.create(
            usuario=usuario,
            fecha=evento_marzo.fecha,
            hora_ingreso=time(8, 5),
            hora_salida=time(10, 0),
            evento=evento_marzo,
            confirmada=True,
        )
        Asistencia.objects.create(
            usuario=usuario,
            fecha=evento_abril.fecha,
            hora_ingreso=time(8, 10),
            hora_salida=time(10, 10),
            evento=evento_abril,
            confirmada=True,
        )

        response = self.client.get(reverse('historial_asistencias'), {
            'fecha_inicio': '2026-03-01',
            'fecha_fin': '2026-04-30',
        })

        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn('Reunion Marzo', content)
        self.assertIn('Reunion Abril', content)

    def test_reporte_filtrado_pdf_indica_varios_eventos_en_rango(self):
        evento_1 = Evento.objects.create(
            nombre='Evento Febrero',
            fecha=date(2026, 2, 10),
            hora_ingreso=time(8, 0),
            activo=False,
        )
        evento_2 = Evento.objects.create(
            nombre='Evento Marzo',
            fecha=date(2026, 3, 5),
            hora_ingreso=time(8, 0),
            activo=False,
        )
        usuario = Usuario.objects.create(nombre='Pablo', apellido='Rios', dni='12345680', estado=Usuario.ESTADO_ACTIVO)

        Asistencia.objects.create(
            usuario=usuario,
            fecha=evento_1.fecha,
            hora_ingreso=time(8, 0),
            hora_salida=time(10, 0),
            evento=evento_1,
            confirmada=True,
        )
        Asistencia.objects.create(
            usuario=usuario,
            fecha=evento_2.fecha,
            hora_ingreso=time(8, 15),
            hora_salida=time(10, 15),
            evento=evento_2,
            confirmada=True,
        )

        response = self.client.get(reverse('descargar_reporte_filtrado_pdf'), {
            'fecha_inicio': '2026-02-01',
            'fecha_fin': '2026-03-31',
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

    def test_filtro_todos_en_rango_incluye_asistencias_y_faltas(self):
        evento_1 = Evento.objects.create(
            nombre='Evento Uno',
            fecha=date(2026, 3, 10),
            hora_ingreso=time(8, 0),
            activo=False,
        )
        evento_2 = Evento.objects.create(
            nombre='Evento Dos',
            fecha=date(2026, 4, 10),
            hora_ingreso=time(8, 0),
            activo=False,
        )
        usuario_asiste = Usuario.objects.create(nombre='Lucia', apellido='Torres', dni='12345681', estado=Usuario.ESTADO_ACTIVO)
        usuario_falta = Usuario.objects.create(nombre='Marco', apellido='Vega', dni='12345682', estado=Usuario.ESTADO_ACTIVO)

        Asistencia.objects.create(
            usuario=usuario_asiste,
            fecha=evento_1.fecha,
            hora_ingreso=time(8, 0),
            hora_salida=time(10, 0),
            evento=evento_1,
            confirmada=True,
        )

        data = get_filtered_attendance_data({
            'fecha_inicio': date(2026, 3, 1),
            'fecha_fin': date(2026, 4, 30),
            'estado': '',
        })

        resumen = sorted(
            (
                item.usuario.dni if hasattr(item, 'usuario') else item.dni,
                getattr(item.evento, 'nombre', None),
                bool(getattr(item, 'is_absent', False)),
            )
            for item in data['unified_report']
        )

        self.assertIn((usuario_asiste.dni, 'Evento Uno', False), resumen)
        self.assertIn((usuario_falta.dni, 'Evento Uno', True), resumen)
        self.assertIn((usuario_asiste.dni, 'Evento Dos', True), resumen)
        self.assertIn((usuario_falta.dni, 'Evento Dos', True), resumen)

    def test_justificacion_rechazada_vuelve_a_falta_en_reportes(self):
        usuario = Usuario.objects.create(
            nombre='Nora',
            apellido='Quispe',
            dni='12345683',
            estado=Usuario.ESTADO_ACTIVO,
        )

        justificacion = Justificacion.objects.create(
            usuario=usuario,
            evento=self.evento,
            motivo='Motivo inicial',
            estado=Justificacion.ESTADO_APROBADO,
            procesado_por=self.user,
        )

        self.assertTrue(
            Asistencia.objects.filter(usuario=usuario, evento=self.evento, es_justificada=True).exists()
        )

        justificacion.estado = Justificacion.ESTADO_RECHAZADO
        justificacion.save()

        self.assertFalse(Asistencia.objects.filter(usuario=usuario, evento=self.evento).exists())

        data = get_filtered_attendance_data({'evento': self.evento, 'estado': 'faltas_injustificadas'})
        resumen = {
            (
                item.usuario.dni if hasattr(item, 'usuario') else item.dni,
                bool(getattr(item, 'is_absent', False)),
                bool(getattr(item, 'es_justificada', False)),
            )
            for item in data['unified_report']
        }
        self.assertIn((usuario.dni, True, False), resumen)

    def test_admin_no_puede_justificar_usuario_exonerado(self):
        usuario_exonerado = Usuario.objects.create(
            nombre='Elio',
            apellido='Exonerado',
            dni='12345684',
            estado=Usuario.ESTADO_EXONERADO,
        )

        response = self.client.post(
            reverse('admin_solicitar_justificacion'),
            data={
                'usuario': usuario_exonerado.id,
                'evento': self.evento.id,
                'motivo': 'Intento de justificación inválido',
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            'Este usuario es exonerado y no se puede proceder con una justificación.',
        )
        self.assertFalse(
            Justificacion.objects.filter(usuario=usuario_exonerado, evento=self.evento).exists()
        )

    def test_api_eventos_con_pendientes_y_regularizacion(self):
        # Crear usuarios y asistencias con solo ingreso (sin salida)
        user1 = Usuario.objects.create(nombre='Carlos', apellido='Gomez', dni='87654321', estado=Usuario.ESTADO_ACTIVO)
        user2 = Usuario.objects.create(nombre='Elena', apellido='Perez', dni='87654322', estado=Usuario.ESTADO_ACTIVO)
        user_exonerado = Usuario.objects.create(nombre='Mario', apellido='Soto', dni='87654323', estado=Usuario.ESTADO_EXONERADO)

        a1 = Asistencia.objects.create(
            usuario=user1,
            evento=self.evento,
            fecha=self.evento.fecha,
            hora_ingreso=time(8, 10),
            hora_salida=None,
            confirmada=False,
        )
        a2 = Asistencia.objects.create(
            usuario=user2,
            evento=self.evento,
            fecha=self.evento.fecha,
            hora_ingreso=time(8, 20),
            hora_salida=None,
            confirmada=False,
        )
        # Asistencia de comunero exonerado en el mismo evento (no debe excluir el evento ni contarse en pendientes)
        Asistencia.objects.create(
            usuario=user_exonerado,
            evento=self.evento,
            fecha=self.evento.fecha,
            hora_ingreso=time(8, 25),
            hora_salida=None,
            confirmada=True,
        )

        # 1. Probar GET api_eventos_con_pendientes
        res_eventos = self.client.get(reverse('api_eventos_con_pendientes'))
        self.assertEqual(res_eventos.status_code, 200)
        data_eventos = res_eventos.json()
        self.assertTrue(data_eventos['success'])
        self.assertEqual(data_eventos['total_pendientes_global'], 2)
        self.assertEqual(len(data_eventos['eventos']), 1)
        self.assertEqual(data_eventos['eventos'][0]['id'], self.evento.id)
        self.assertEqual(data_eventos['eventos'][0]['total_pendientes'], 2)

        # 2. Probar GET api_evento_pendientes
        res_pendientes = self.client.get(reverse('api_evento_pendientes', args=[self.evento.id]))
        self.assertEqual(res_pendientes.status_code, 200)
        data_pendientes = res_pendientes.json()
        self.assertTrue(data_pendientes['success'])
        self.assertEqual(data_pendientes['total'], 2)
        p_ids = [p['asistencia_id'] for p in data_pendientes['pendientes']]
        self.assertIn(a1.id, p_ids)
        self.assertIn(a2.id, p_ids)

        # 3. Probar POST regularizar_pendientes_evento (regularizar solo user1 primero)
        res_reg = self.client.post(
            reverse('regularizar_pendientes_evento'),
            data={
                'evento_id': self.evento.id,
                'marcar_todos': False,
                'asistencia_ids': [a1.id],
                'hora_salida': '12:30',
                'observacion': 'Salida por otra puerta regularizada',
            },
            content_type='application/json',
        )
        self.assertEqual(res_reg.status_code, 200)
        data_reg = res_reg.json()
        self.assertTrue(data_reg['success'])
        self.assertEqual(data_reg['count'], 1)

        # Verificar a1 actualizado
        a1.refresh_from_db()
        self.assertTrue(a1.confirmada)
        self.assertEqual(a1.hora_salida, time(12, 30))

        # Verificar log en LogAccion
        log_entry = LogAccion.objects.filter(accion='Regularización de Asistencias Pendientes').first()
        self.assertIsNotNone(log_entry)
        self.assertIn(self.evento.nombre, log_entry.descripcion)

        # 4. Probar POST regularizar con marcar_todos=True para regularizar a user2
        res_reg_todos = self.client.post(
            reverse('regularizar_pendientes_evento'),
            data={
                'evento_id': self.evento.id,
                'marcar_todos': True,
                'hora_salida': '13:00',
            },
            content_type='application/json',
        )
        self.assertEqual(res_reg_todos.status_code, 200)
        self.assertEqual(res_reg_todos.json()['count'], 1)

        a2.refresh_from_db()
        self.assertTrue(a2.confirmada)
        self.assertEqual(a2.hora_salida, time(13, 0))

        # 5. Verificar que ahora el evento ya no tiene pendientes
        res_vacio = self.client.get(reverse('api_eventos_con_pendientes'))
        self.assertEqual(res_vacio.json()['total_pendientes_global'], 0)
        self.assertEqual(len(res_vacio.json()['eventos']), 0)

        # 6. Test de permisos: usuario sin 'can_manage_users' es denegado
        user_sin_permiso = User.objects.create_user(username='regular_user', password='password')
        client_sin_permiso = Client()
        client_sin_permiso.force_login(user_sin_permiso)
        res_unauthorized = client_sin_permiso.get(reverse('api_eventos_con_pendientes'))
        self.assertEqual(res_unauthorized.status_code, 403)

        res_unauthorized_post = client_sin_permiso.post(
            reverse('regularizar_pendientes_evento'),
            data={'evento_id': self.evento.id, 'marcar_todos': True},
            content_type='application/json',
        )
        self.assertEqual(res_unauthorized_post.status_code, 403)

        # 7. Test de validación: POST sin IDs y con marcar_todos=False retorna 400
        res_bad_request = self.client.post(
            reverse('regularizar_pendientes_evento'),
            data={'evento_id': self.evento.id, 'marcar_todos': False, 'asistencia_ids': []},
            content_type='application/json',
        )
        self.assertEqual(res_bad_request.status_code, 400)


class AdminMassPhotoUploadTests(TestCase):
    def setUp(self):
        self.temp_media_root = tempfile.mkdtemp(prefix='test-media-')
        self._media_override = override_settings(MEDIA_ROOT=self.temp_media_root)
        self._media_override.enable()

        self.admin_user = User.objects.create_superuser(
            username='admin',
            email='admin@example.com',
            password='password',
        )
        self.client = Client()
        self.client.force_login(self.admin_user)

    def tearDown(self):
        self._media_override.disable()
        shutil.rmtree(self.temp_media_root, ignore_errors=True)

    def _build_image(self, name, color):
        buffer = BytesIO()
        Image.new('RGB', (320, 320), color=color).save(buffer, format='JPEG')
        return SimpleUploadedFile(name, buffer.getvalue(), content_type='image/jpeg')

    def test_carga_masiva_admin_ajax_procesa_archivos(self):
        usuario = Usuario.objects.create(
            nombre='Julia',
            apellido='Ramos',
            dni='12345678',
            estado=Usuario.ESTADO_ACTIVO,
        )

        response = self.client.post(
            reverse('admin:asistencia_usuario_carga_masiva_fotos'),
            {'fotos': [self._build_image('12345678.jpg', (12, 120, 200))]},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )

        self.assertEqual(response.status_code, 200, response.content.decode())
        payload = response.json()
        self.assertTrue(payload['ok'])
        self.assertEqual(payload['actualizados'], 1)

        usuario.refresh_from_db()
        self.assertTrue(bool(usuario.foto_perfil))

    def test_carga_masiva_reemplaza_foto_anterior_sin_dejar_huerfanos(self):
        usuario = Usuario.objects.create(
            nombre='Julia',
            apellido='Ramos',
            dni='12345678',
            estado=Usuario.ESTADO_ACTIVO,
        )

        first_response = self.client.post(
            reverse('admin:asistencia_usuario_carga_masiva_fotos'),
            {'fotos': [self._build_image('12345678.jpg', (12, 120, 200))]},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(first_response.status_code, 200, first_response.content.decode())

        second_response = self.client.post(
            reverse('admin:asistencia_usuario_carga_masiva_fotos'),
            {'fotos': [self._build_image('12345678.jpg', (220, 80, 40))]},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(second_response.status_code, 200, second_response.content.decode())

        usuario.refresh_from_db()
        self.assertEqual(usuario.foto_perfil.name, 'perfil_fotos/perfil_12345678.jpg')

        photo_dir = os.path.join(settings.MEDIA_ROOT, 'perfil_fotos')
        self.assertEqual(sorted(os.listdir(photo_dir)), ['perfil_12345678.jpg'])


class CarnetsZipDownloadTests(TestCase):
    def setUp(self):
        self.temp_media_root = tempfile.mkdtemp(prefix='test-media-')
        self._media_override = override_settings(MEDIA_ROOT=self.temp_media_root)
        self._media_override.enable()

        self.user = User.objects.create_user(username='manager_zip', password='password')
        self.user.user_permissions.add(Permission.objects.get(codename='can_manage_users'))
        self.client = Client()
        self.client.force_login(self.user)

        ConfiguracionSistema.objects.create(nombre_institucion='QUIULACOCHA', tolerancia_minutos=15)

    def tearDown(self):
        self._media_override.disable()
        shutil.rmtree(self.temp_media_root, ignore_errors=True)

    def _create_photo_file(self, color=(24, 120, 196)):
        source = BytesIO()
        Image.new('RGB', (500, 500), color=color).save(source, format='PNG')
        return SimpleUploadedFile('avatar.png', source.getvalue(), content_type='image/png')

    def _build_fake_pdf_modules(self):
        class FakeHTML:
            def __init__(self, string):
                self.string = string

            def write_pdf(self):
                normalized = self.string.lower()
                if 'usuarios sin foto' in normalized or 'pendientes de fotograf' in normalized:
                    return b'INCIDENTES_PDF'
                if 'no se encontraron usuarios' in normalized:
                    return b'PDF_VACIO'
                return b'CARNETS_PDF'

        class FakePdfReader:
            def __init__(self, buffer):
                self.pages = [buffer.getvalue()]

        class FakePdfWriter:
            def __init__(self):
                self.pages = []

            def add_page(self, page):
                self.pages.append(page)

            def write(self, buffer):
                buffer.write(b''.join(self.pages) or b'PDF_VACIO')

        return (
            types.SimpleNamespace(HTML=FakeHTML),
            types.SimpleNamespace(PdfReader=FakePdfReader, PdfWriter=FakePdfWriter),
        )

    def test_descarga_zip_mixta_incluye_tres_archivos_esperados(self):
        Usuario.objects.create(
            nombre='Ana',
            apellido='ConFoto',
            dni='12345690',
            estado=Usuario.ESTADO_ACTIVO,
            foto_perfil=self._create_photo_file(),
        )
        Usuario.objects.create(
            nombre='Luis',
            apellido='SinFoto',
            dni='12345691',
            estado=Usuario.ESTADO_PASIVO,
        )

        fake_weasyprint, fake_pypdf = self._build_fake_pdf_modules()
        with mock.patch.dict(sys.modules, {'weasyprint': fake_weasyprint, 'pypdf': fake_pypdf}):
            response = self.client.get(reverse('descargar_todos_carnets_pdf'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/zip')
        self.assertIn('attachment;', response['Content-Disposition'])
        self.assertIn('carnets_y_pendientes_', response['Content-Disposition'])

        with zipfile.ZipFile(BytesIO(response.content)) as zip_file:
            self.assertEqual(
                sorted(zip_file.namelist()),
                ['carnets_con_foto.pdf', 'carnets_sin_foto.pdf', 'reporte_usuarios_sin_foto.pdf'],
            )
            self.assertEqual(zip_file.read('carnets_con_foto.pdf'), b'CARNETS_PDF')
            self.assertEqual(zip_file.read('carnets_sin_foto.pdf'), b'CARNETS_PDF')
            self.assertEqual(zip_file.read('reporte_usuarios_sin_foto.pdf'), b'INCIDENTES_PDF')

    def test_descarga_zip_solo_con_foto_genera_placeholders_restantes(self):
        Usuario.objects.create(
            nombre='Rosa',
            apellido='Valida',
            dni='12345692',
            estado=Usuario.ESTADO_ACTIVO,
            foto_perfil=self._create_photo_file(color=(0, 180, 90)),
        )

        fake_weasyprint, fake_pypdf = self._build_fake_pdf_modules()
        with mock.patch.dict(sys.modules, {'weasyprint': fake_weasyprint, 'pypdf': fake_pypdf}):
            response = self.client.get(reverse('descargar_todos_carnets_pdf'))

        with zipfile.ZipFile(BytesIO(response.content)) as zip_file:
            self.assertEqual(
                sorted(zip_file.namelist()),
                ['carnets_con_foto.pdf', 'carnets_sin_foto.pdf', 'reporte_usuarios_sin_foto.pdf'],
            )

    def test_descarga_zip_reporta_archivo_foto_faltante_como_incidencia(self):
        usuario = Usuario.objects.create(
            nombre='Mario',
            apellido='ArchivoPerdido',
            dni='12345693',
            estado=Usuario.ESTADO_ACTIVO,
            foto_perfil=self._create_photo_file(color=(180, 60, 60)),
        )
        os.remove(usuario.foto_perfil.path)

        fake_weasyprint, fake_pypdf = self._build_fake_pdf_modules()
        with mock.patch.dict(sys.modules, {'weasyprint': fake_weasyprint, 'pypdf': fake_pypdf}):
            response = self.client.get(reverse('descargar_todos_carnets_pdf'))

        with zipfile.ZipFile(BytesIO(response.content)) as zip_file:
            self.assertEqual(
                sorted(zip_file.namelist()),
                ['carnets_sin_foto.pdf', 'reporte_usuarios_sin_foto.pdf'],
            )

    @mock.patch('asistencia.views._launch_carnets_generation_job')
    def test_solicitar_descarga_carnets_crea_solicitud_pendiente(self, launch_job):
        response = self.client.post(reverse('solicitar_descarga_carnets'))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload['ok'])

        solicitud = SolicitudDescargaCarnets.objects.get()
        self.assertEqual(solicitud.solicitado_por, self.user)
        self.assertEqual(solicitud.estado, SolicitudDescargaCarnets.ESTADO_PENDIENTE)
        launch_job.assert_called_once_with(solicitud.id)

    def test_solicitar_descarga_carnets_bloquea_si_hay_otra_activa(self):
        SolicitudDescargaCarnets.objects.create(
            solicitado_por=self.user,
            estado=SolicitudDescargaCarnets.ESTADO_PROCESANDO,
        )

        response = self.client.post(reverse('solicitar_descarga_carnets'))

        self.assertEqual(response.status_code, 409)
        payload = response.json()
        self.assertFalse(payload['ok'])

    def test_descargar_archivo_carnets_listo_devuelve_zip(self):
        solicitud = SolicitudDescargaCarnets.objects.create(
            solicitado_por=self.user,
            estado=SolicitudDescargaCarnets.ESTADO_LISTO,
            nombre_archivo='carnets_demo.zip',
        )
        solicitud.archivo_zip.save('carnets_demo.zip', ContentFile(b'ZIPDATA'), save=True)

        response = self.client.get(reverse('descargar_archivo_carnets', args=[solicitud.id]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/zip')


class GestionCarnetsModuleTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='manager_carnets', password='password')
        self.user.user_permissions.add(Permission.objects.get(codename='can_manage_users'))
        self.client = Client()
        self.client.force_login(self.user)
        self.url = reverse('modulo_gestion_carnets')

    def _create_photo_file(self, name='foto.jpg', color=(40, 100, 180)):
        image = Image.new('RGB', (50, 50), color=color)
        buffer = BytesIO()
        image.save(buffer, format='JPEG')
        return SimpleUploadedFile(name, buffer.getvalue(), content_type='image/jpeg')

    def test_reposicion_perdida_guarda_motivo_estandar_y_auditoria(self):
        usuario = Usuario.objects.create(
            nombre='Lucia',
            apellido='Prueba',
            dni='76543210',
            estado=Usuario.ESTADO_ACTIVO,
            foto_perfil=self._create_photo_file(),
        )
        version_inicial = usuario.qr_version
        HistorialCarnet.objects.create(
            usuario=usuario,
            motivo='Primer Carnet',
            estado='Activo',
            entregado_por=self.user,
        )

        response = self.client.post(
            self.url,
            {'usuario_id': usuario.id, 'accion': 'reposicion_perdida'},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(
            HistorialCarnet.objects.filter(
                usuario=usuario,
                motivo='Reposición por Pérdida/Robo',
                estado='Activo',
                entregado_por=self.user,
            ).exists()
        )
        self.assertTrue(
            LogAccion.objects.filter(
                usuario=self.user,
                accion='Emisión de carnet físico',
                descripcion__contains='Reposición por Pérdida/Robo',
            ).exists()
        )
        usuario.refresh_from_db()
        self.assertEqual(usuario.qr_version, version_inicial + 1)

    def test_primer_carnet_requiere_foto(self):
        usuario = Usuario.objects.create(
            nombre='Mario',
            apellido='SinFoto',
            dni='76543211',
            estado=Usuario.ESTADO_ACTIVO,
        )

        response = self.client.post(
            self.url,
            {'usuario_id': usuario.id, 'accion': 'primer_carnet'},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(HistorialCarnet.objects.filter(usuario=usuario).exists())
        mensajes = [str(m) for m in response.context['messages']]
        self.assertTrue(any('sin fotografía' in mensaje for mensaje in mensajes))

    def test_primer_carnet_no_permite_duplicar_activo(self):
        usuario = Usuario.objects.create(
            nombre='Ana',
            apellido='Activa',
            dni='76543212',
            estado=Usuario.ESTADO_ACTIVO,
            foto_perfil=self._create_photo_file(name='activa.jpg', color=(120, 120, 20)),
        )
        HistorialCarnet.objects.create(
            usuario=usuario,
            motivo='Primer Carnet',
            estado='Activo',
            entregado_por=self.user,
        )

        response = self.client.post(
            self.url,
            {'usuario_id': usuario.id, 'accion': 'primer_carnet'},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(HistorialCarnet.objects.filter(usuario=usuario, estado='Activo').count(), 1)
        mensajes = [str(m) for m in response.context['messages']]
        self.assertTrue(any('ya cuenta con un carnet activo' in mensaje for mensaje in mensajes))


class QRSecurityAndRescueModeTests(TestCase):
    def setUp(self):
        self.config, _ = ConfiguracionSistema.objects.get_or_create(
            defaults={'nombre_institucion': 'QUIULACOCHA', 'tolerancia_minutos': 15}
        )

    def test_qr_encode_and_decode_roundtrip(self):
        import uuid
        from .qr_security import build_qr_encoded_payload, decode_qr_encoded_payload

        qr_uid = uuid.uuid4()
        payload = build_qr_encoded_payload(qr_uid, 1)

        result = decode_qr_encoded_payload(payload)
        self.assertEqual(result['kind'], 'uid_v1')
        self.assertEqual(result['qr_uid'], qr_uid)
        self.assertEqual(result['qr_version'], 1)

    def test_qr_rescue_mode_accepts_alien_signature(self):
        import uuid
        import base64
        from django.core.signing import Signer
        from .qr_security import decode_qr_encoded_payload, QR_SIGNING_SALT

        self.config.permitir_carnets_antiguos_sin_firma = True
        self.config.save()

        qr_uid = uuid.uuid4()
        raw = f"v1|{qr_uid}|1"
        alien_signer = Signer(key="alien-secret-from-old-server", salt=QR_SIGNING_SALT)
        alien_signed = alien_signer.sign(raw)
        encoded_alien_qr = base64.b64encode(alien_signed.encode()).decode()

        result = decode_qr_encoded_payload(encoded_alien_qr)
        self.assertEqual(result['kind'], 'uid_v1')
        self.assertEqual(result['qr_uid'], qr_uid)
        self.assertEqual(result['qr_version'], 1)

    def test_qr_rescue_mode_disabled_rejects_alien_signature(self):
        import uuid
        import base64
        from django.core.signing import BadSignature, Signer
        from .qr_security import decode_qr_encoded_payload, QR_SIGNING_SALT

        self.config.permitir_carnets_antiguos_sin_firma = False
        self.config.save()

        qr_uid = uuid.uuid4()
        raw = f"v1|{qr_uid}|1"
        alien_signer = Signer(key="alien-secret-from-old-server", salt=QR_SIGNING_SALT)
        alien_signed = alien_signer.sign(raw)
        encoded_alien_qr = base64.b64encode(alien_signed.encode()).decode()

        with self.assertRaises(BadSignature):
            decode_qr_encoded_payload(encoded_alien_qr)

    def test_qr_config_signing_key_and_fallback(self):
        import uuid
        import base64
        from django.core.signing import Signer
        from .qr_security import decode_qr_encoded_payload, QR_SIGNING_SALT

        self.config.permitir_carnets_antiguos_sin_firma = False
        self.config.qr_signing_key = "current-custom-key"
        self.config.qr_fallback_keys = "fallback-key-1, fallback-key-2"
        self.config.save()

        qr_uid1 = uuid.uuid4()
        signer_current = Signer(key="current-custom-key", salt=QR_SIGNING_SALT)
        encoded1 = base64.b64encode(signer_current.sign(f"v1|{qr_uid1}|1").encode()).decode()
        res1 = decode_qr_encoded_payload(encoded1)
        self.assertEqual(res1['qr_uid'], qr_uid1)

        qr_uid2 = uuid.uuid4()
        signer_fallback = Signer(key="fallback-key-2", salt=QR_SIGNING_SALT)
        encoded2 = base64.b64encode(signer_fallback.sign(f"v1|{qr_uid2}|2").encode()).decode()
        res2 = decode_qr_encoded_payload(encoded2)
        self.assertEqual(res2['qr_uid'], qr_uid2)

    def test_qr_malformed_input_raises_bad_signature(self):
        from django.core.signing import BadSignature
        from .qr_security import decode_qr_encoded_payload

        with self.assertRaises(BadSignature):
            decode_qr_encoded_payload("not-valid-base64-!@#$%^&*()")

        with self.assertRaises(BadSignature):
            import base64
            # valid base64 but invalid content
            decode_qr_encoded_payload(base64.b64encode(b"v1|not-a-uuid|1").decode())


class BackupRestoreWebTests(TransactionTestCase):
    def setUp(self):
        self.temp_media_root = tempfile.mkdtemp(prefix='test-media-')
        self._media_override = override_settings(MEDIA_ROOT=self.temp_media_root)
        self._media_override.enable()

        self.superuser = User.objects.create_superuser(
            username='admin_backup',
            email='admin@example.com',
            password='password'
        )
        self.regular_user = User.objects.create_user(
            username='regular',
            password='password'
        )
        self.regular_user.user_permissions.add(Permission.objects.get(codename='can_manage_users'))

    def tearDown(self):
        self._media_override.disable()
        shutil.rmtree(self.temp_media_root, ignore_errors=True)

    def test_non_superuser_denied(self):
        self.client.force_login(self.regular_user)
        response = self.client.post(reverse('restaurar_backup'), follow=True)
        self.assertEqual(response.status_code, 200)
        messages = [str(m) for m in response.context['messages']]
        self.assertTrue(any('Solo superusuarios pueden restaurar' in m for m in messages))

    def test_invalid_extension_rejected(self):
        self.client.force_login(self.superuser)
        bad_file = SimpleUploadedFile("backup.zip", b"fake zip content", content_type="application/zip")
        response = self.client.post(reverse('restaurar_backup'), {'backup_file': bad_file}, follow=True)
        self.assertEqual(response.status_code, 200)
        messages = [str(m) for m in response.context['messages']]
        self.assertTrue(any('debe ser un .tar.gz valido' in m for m in messages))

    def test_restore_backup_tar_gz_with_media_and_database(self):
        import tarfile
        self.client.force_login(self.superuser)

        test_tar_dir = tempfile.mkdtemp()
        try:
            media_content_dir = os.path.join(test_tar_dir, 'media_content')
            os.makedirs(os.path.join(media_content_dir, 'perfil_fotos'), exist_ok=True)
            with open(os.path.join(media_content_dir, 'perfil_fotos', 'test_photo.jpg'), 'wb') as f:
                f.write(b'fake_photo_data')

            media_tar_path = os.path.join(test_tar_dir, 'media.tar.gz')
            with tarfile.open(media_tar_path, 'w:gz') as mt:
                mt.add(media_content_dir, arcname='media')

            db_sql_path = os.path.join(test_tar_dir, 'database.sql')
            with open(db_sql_path, 'w', encoding='utf-8') as f:
                f.write('-- empty sql dump\n')

            main_tar_path = os.path.join(test_tar_dir, 'backup.tar.gz')
            with tarfile.open(main_tar_path, 'w:gz') as main_tar:
                main_tar.add(media_tar_path, arcname='media.tar.gz')
                main_tar.add(db_sql_path, arcname='database.sql')

            with open(main_tar_path, 'rb') as f:
                uploaded = SimpleUploadedFile('backup_test.tar.gz', f.read(), content_type='application/gzip')

            response = self.client.post(reverse('restaurar_backup'), {'backup_file': uploaded}, follow=True)
            self.assertEqual(response.status_code, 200)

            restored_photo = os.path.join(settings.MEDIA_ROOT, 'perfil_fotos', 'test_photo.jpg')
            self.assertTrue(os.path.exists(restored_photo))
        finally:
            shutil.rmtree(test_tar_dir, ignore_errors=True)

    def test_restore_backup_tar_gz_ajax(self):
        import tarfile
        self.client.force_login(self.superuser)

        test_tar_dir = tempfile.mkdtemp()
        try:
            media_content_dir = os.path.join(test_tar_dir, 'media_content')
            os.makedirs(os.path.join(media_content_dir, 'perfil_fotos'), exist_ok=True)
            with open(os.path.join(media_content_dir, 'perfil_fotos', 'ajax_photo.jpg'), 'wb') as f:
                f.write(b'fake_ajax_photo')

            media_tar_path = os.path.join(test_tar_dir, 'media.tar.gz')
            with tarfile.open(media_tar_path, 'w:gz') as mt:
                mt.add(media_content_dir, arcname='media')

            db_sql_path = os.path.join(test_tar_dir, 'database.sql')
            with open(db_sql_path, 'w', encoding='utf-8') as f:
                f.write('-- empty sql dump\n')

            main_tar_path = os.path.join(test_tar_dir, 'backup.tar.gz')
            with tarfile.open(main_tar_path, 'w:gz') as main_tar:
                main_tar.add(media_tar_path, arcname='media.tar.gz')
                main_tar.add(db_sql_path, arcname='database.sql')

            with open(main_tar_path, 'rb') as f:
                uploaded = SimpleUploadedFile('backup_test.tar.gz', f.read(), content_type='application/gzip')

            response = self.client.post(
                reverse('restaurar_backup'),
                {'backup_file': uploaded},
                HTTP_X_REQUESTED_WITH='XMLHttpRequest',
            )
            self.assertEqual(response.status_code, 200)
            data = response.json()
            self.assertTrue(data.get('ok'))
            self.assertIn('Sistema restaurado exitosamente', data.get('message'))
        finally:
            shutil.rmtree(test_tar_dir, ignore_errors=True)



